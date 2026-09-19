"""Bundled skill scripts emit their JSON result as UTF-8, whatever the console code page.

``print()`` encodes through ``sys.stdout.encoding``. On Windows that is the
console code page (cp1252 here, cp936/cp932 on CJK systems), not UTF-8, so a
result carrying a character outside that page raised ``UnicodeEncodeError``
and the script died with a traceback instead of returning anything (#2334).
The ``--out`` branch of the same scripts was never affected because it passes
``encoding="utf-8"`` explicitly, which is what made the stdout path the odd
one out rather than a platform limitation.

The code page is **simulated**, not skipped: ``sys.stdout`` is replaced with a
cp1252-encoded text stream over a ``BytesIO``. A ``skipif(sys.platform !=
"win32")`` would run on only half of CI, and the defect is in the encoder
choice, not in Windows. The same substitution reproduces it on any host.

Three cases here pass on an unfixed tree by design and say so in their own
docstring: their payload is integers only, so it cannot carry a character the
code page rejects today. They are guards that the write path stays UTF-8 when
a string field is added to those payloads later.
"""

from __future__ import annotations

import importlib.util
import io
import json
import sys
from pathlib import Path
from typing import Any

import pytest

ROOT = Path(__file__).resolve().parents[1]
BUNDLED = ROOT / "src" / "agentos" / "skills" / "bundled"

#: Characters outside cp1252: CJK, and a Vietnamese vowel with two diacritics.
NON_ASCII = "日本語のテキスト cổ phiếu"


def _load(relative: str, name: str) -> Any:
    """Import a bundled script by path, the way the other skill tests do."""
    spec = importlib.util.spec_from_file_location(name, BUNDLED / relative)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    # Registered before exec: a dataclass in the module resolves its own
    # annotations through ``sys.modules[cls.__module__]``, which is absent for
    # a module built straight from a spec.
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


class CodePageStdout:
    """A ``sys.stdout`` whose text layer only accepts one legacy code page.

    Installed as a context manager **inside** the test body rather than from a
    fixture: pytest re-activates its own capture at the start of each test
    phase, so a ``sys.stdout`` swapped in during fixture setup is replaced
    again before the test runs.
    """

    def __init__(self, encoding: str = "cp1252") -> None:
        self.sink = io.BytesIO()
        self.stream = io.TextIOWrapper(self.sink, encoding=encoding, newline="")
        self._saved: Any = None

    def __enter__(self) -> CodePageStdout:
        self._saved = sys.stdout
        sys.stdout = self.stream
        return self

    def __exit__(self, *_exc: Any) -> None:
        sys.stdout = self._saved
        self.stream.flush()
        # Detach so the wrapper's finalizer cannot close the sink we still read.
        self.stream.detach()

    def text(self) -> str:
        return self.sink.getvalue().decode("utf-8")

    def payload(self) -> Any:
        return json.loads(self.text())


def _argv(monkeypatch: pytest.MonkeyPatch, *args: str) -> None:
    monkeypatch.setattr(sys, "argv", list(args))


# ── docx ────────────────────────────────────────────────────────────────────


def _make_docx(path: Path, text: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))


def test_inspect_docx_emits_utf8_on_a_code_page_console(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    inspect_docx = _load("docx/scripts/inspect_docx.py", "inspect_docx")
    source = tmp_path / "cjk.docx"
    _make_docx(source, NON_ASCII)
    _argv(monkeypatch, "inspect_docx.py", str(source))

    with CodePageStdout() as code_page_stdout:
        assert inspect_docx.main() == 0
    assert NON_ASCII in json.dumps(code_page_stdout.payload(), ensure_ascii=False)


def test_edit_docx_result_is_written_as_utf8(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Guard — passes either way today: the payload is ``{"applied": <int>}``.

    Pinned so the write path is already UTF-8 the day that payload grows a
    string field (a sheet name, a path, an error) and the defect would
    otherwise come back silently.
    """
    edit_docx = _load("docx/scripts/edit_docx.py", "edit_docx")
    source = tmp_path / "in.docx"
    _make_docx(source, NON_ASCII)
    ops = tmp_path / "ops.json"
    ops.write_text(
        json.dumps([{"op": "replace_text", "find": NON_ASCII, "replace": "x"}]),
        encoding="utf-8",
    )
    out = tmp_path / "out.docx"
    _argv(monkeypatch, "edit_docx.py", str(source), str(ops), "--out", str(out))

    with CodePageStdout() as code_page_stdout:
        assert edit_docx.main() == 0
    assert code_page_stdout.text().endswith("\n")
    assert "applied" in code_page_stdout.payload()


# ── pdf-toolkit ─────────────────────────────────────────────────────────────


def _make_pdf(path: Path, title: str) -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_metadata({"/Title": title})
    with path.open("wb") as handle:
        writer.write(handle)


def _make_form(path: Path) -> None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path), pagesize=LETTER)
    pdf.acroForm.textfield(name="full_name", x=72, y=700, width=200, height=20)
    pdf.showPage()
    pdf.save()


def test_pdf_extract_emits_utf8_on_a_code_page_console(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    extract = _load("pdf-toolkit/scripts/extract.py", "extract")
    source = tmp_path / "cjk.pdf"
    _make_pdf(source, NON_ASCII)
    _argv(monkeypatch, "extract.py", str(source))

    with CodePageStdout() as code_page_stdout:
        assert extract.main() == 0
    assert NON_ASCII in json.dumps(code_page_stdout.payload(), ensure_ascii=False)


def test_form_fill_list_fields_emits_utf8_on_a_code_page_console(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    form_fill = _load("pdf-toolkit/scripts/form_fill.py", "form_fill")
    form = tmp_path / "form.pdf"
    _make_form(form)
    filled = tmp_path / "filled.pdf"
    form_fill.fill(form, {"full_name": NON_ASCII}, filled)
    _argv(monkeypatch, "form_fill.py", str(filled), "--list-fields")

    with CodePageStdout() as code_page_stdout:
        assert form_fill.main() == 0
    assert NON_ASCII in code_page_stdout.text()


def test_form_fill_result_is_written_as_utf8(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Guard — passes either way today: the payload is two integers."""
    form_fill = _load("pdf-toolkit/scripts/form_fill.py", "form_fill")
    form = tmp_path / "form.pdf"
    _make_form(form)
    data = tmp_path / "data.json"
    data.write_text(json.dumps({"full_name": NON_ASCII}), encoding="utf-8")
    out = tmp_path / "filled.pdf"
    _argv(monkeypatch, "form_fill.py", str(form), str(data), "--out", str(out))

    with CodePageStdout() as code_page_stdout:
        assert form_fill.main() == 0
    assert code_page_stdout.payload()["fields"] == 1


def test_pdf_merge_emits_a_non_ascii_output_path_as_utf8(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    merge = _load("pdf-toolkit/scripts/merge.py", "merge")
    first = tmp_path / "a.pdf"
    second = tmp_path / "b.pdf"
    _make_pdf(first, "a")
    _make_pdf(second, "b")
    out = tmp_path / f"{NON_ASCII}.pdf"
    _argv(monkeypatch, "merge.py", str(first), str(second), "--out", str(out))

    with CodePageStdout() as code_page_stdout:
        assert merge.main() == 0
    assert NON_ASCII in code_page_stdout.payload()["out"]


def test_pdf_split_emits_non_ascii_output_paths_as_utf8(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    split = _load("pdf-toolkit/scripts/split.py", "split")
    source = tmp_path / "in.pdf"
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    with source.open("wb") as handle:
        writer.write(handle)
    out_dir = tmp_path / NON_ASCII
    _argv(monkeypatch, "split.py", str(source), "--pages", "1,2", "--out", str(out_dir))

    with CodePageStdout() as code_page_stdout:
        assert split.main() == 0
    assert any(NON_ASCII in name for name in code_page_stdout.payload()["files"])


# ── xlsx ────────────────────────────────────────────────────────────────────


def test_edit_xlsx_result_is_written_as_utf8(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Guard — passes either way today: the payload is ``{"applied": <int>}``."""
    edit_xlsx = _load("xlsx/scripts/edit_xlsx.py", "edit_xlsx")
    from openpyxl import Workbook

    source = tmp_path / "in.xlsx"
    book = Workbook()
    book.active.title = "Sheet1"
    book.save(str(source))
    ops = tmp_path / "ops.json"
    ops.write_text(
        json.dumps([{"op": "set_cell", "sheet": "Sheet1", "row": 1, "col": 1, "value": NON_ASCII}]),
        encoding="utf-8",
    )
    out = tmp_path / "out.xlsx"
    _argv(monkeypatch, "edit_xlsx.py", str(source), str(ops), "--out", str(out))

    with CodePageStdout() as code_page_stdout:
        assert edit_xlsx.main() == 0
    assert code_page_stdout.payload()["applied"] == 1


# ── robinhood ───────────────────────────────────────────────────────────────


def test_chain_stocks_invalid_rpc_url_error_is_utf8(monkeypatch: pytest.MonkeyPatch) -> None:
    chain_stocks = _load("robinhood-chain-stocks/scripts/chain_stocks.py", "chain_stocks")
    bad = f"ftp://{NON_ASCII}.test"
    _argv(monkeypatch, "chain_stocks.py", "--query", "AAPL", "--rpc-url", bad, "--no-cards")

    with CodePageStdout() as code_page_stdout:
        assert chain_stocks.main() == 0
    assert NON_ASCII in code_page_stdout.payload()["error"]


def test_chain_stocks_resolve_failure_echoes_the_query_as_utf8(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    chain_stocks = _load("robinhood-chain-stocks/scripts/chain_stocks.py", "chain_stocks")

    def _boom(*_args: Any, **_kwargs: Any) -> Any:
        raise ValueError("no match")

    monkeypatch.setattr(chain_stocks, "_resolve_target", _boom)
    _argv(monkeypatch, "chain_stocks.py", "--query", NON_ASCII, "--no-cards")

    with CodePageStdout() as code_page_stdout:
        assert chain_stocks.main() == 0
    assert code_page_stdout.payload()["query"] == NON_ASCII


def test_chain_stocks_result_echoes_the_query_as_utf8(monkeypatch: pytest.MonkeyPatch) -> None:
    chain_stocks = _load("robinhood-chain-stocks/scripts/chain_stocks.py", "chain_stocks")
    address = "0x" + "ab" * 20

    monkeypatch.setattr(
        chain_stocks,
        "_resolve_target",
        lambda *_a, **_k: (address, {"name": NON_ASCII, "symbol": "AAPL"}, []),
    )
    monkeypatch.setattr(chain_stocks, "inspect_token", lambda *_a, **_k: {"isStockToken": True})
    _argv(monkeypatch, "chain_stocks.py", "--query", NON_ASCII, "--no-cards", "--no-price")

    with CodePageStdout() as code_page_stdout:
        assert chain_stocks.main() == 0
    assert code_page_stdout.payload()["query"] == NON_ASCII


def test_rwa_lookup_result_echoes_the_query_as_utf8(monkeypatch: pytest.MonkeyPatch) -> None:
    rwa_lookup = _load("robinhood-rwa-addresses/scripts/rwa_lookup.py", "rwa_lookup")

    monkeypatch.setattr(rwa_lookup, "_fetch_tokens", lambda *_a, **_k: [])
    _argv(monkeypatch, "rwa_lookup.py", "--query", NON_ASCII, "--no-verify", "--no-cards")

    with CodePageStdout() as code_page_stdout:
        assert rwa_lookup.main() == 0
    assert code_page_stdout.payload()["query"] == NON_ASCII


def test_text_file_read_emits_utf8_on_a_code_page_console(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    text_file_read = _load("text-file-read/scripts/read.py", "text_file_read")
    sample_file = tmp_path / "sample.txt"
    sample_file.write_text(NON_ASCII, encoding="utf-8")
    _argv(monkeypatch, "read.py", "--input", str(sample_file))

    with CodePageStdout() as code_page_stdout:
        assert text_file_read.main() == 0
    assert code_page_stdout.text() == NON_ASCII
